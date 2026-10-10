"""
Тестове за Фаза 13: един разбор на текста на отчета (report_text) и износ на запазен отчет в DOCX и Markdown.

Износът е само за собственика, без AI и без такса, с безопасно име на файла и без сурови знаци за форматиране.
Пускане: python -m unittest discover -s tests -p "test_report_export.py"
"""
import io
import unittest
import zipfile
from unittest import mock

import testenv  # noqa: F401  (трябва да е преди main)
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import data_api  # noqa: E402
import main  # noqa: E402
import report_text  # noqa: E402
from rate_limit import limiter  # noqa: E402
from testenv import CHART, register_and_login  # noqa: E402

MIXED = ("<h2>Общ преглед на периода</h2>\n**Бележка.** Текст.\n\n"
         "<h2>Октомври 2026</h2>\n<p>Първи абзац с <strong>удебелено</strong> и *курсив*.</p>\n"
         "## Подзаглавие\n- първа точка\n- втора точка\n1. едно\n2. две\n━━━━━━━━\nКрай.")


class ParseTest(unittest.TestCase):
    def kinds(self, text):
        return [(b.kind, b.level) for b in report_text.parse(text)]

    def test_markdown_blocks(self):
        text = "# Заглавие\n\nАбзац един\nпродължава.\n\n## Раздел\n- а\n- б\n\n1. x\n2) y\n\n---\n\nКрай"
        self.assertEqual(self.kinds(text), [("h", 1), ("p", 0), ("h", 2), ("ul", 0), ("ol", 0), ("hr", 0), ("p", 0)])
        self.assertEqual(report_text.parse(text)[1].items, ["Абзац един продължава."])
        self.assertEqual(report_text.parse(text)[3].items, ["а", "б"])

    def test_html_from_the_app_becomes_the_same_blocks(self):
        blocks = report_text.parse("<h2>Месец</h2><p>Текст <b>важен</b></p><ul><li>един</li><li>два</li></ul>")
        self.assertEqual([(b.kind, b.level) for b in blocks], [("h", 2), ("p", 0), ("ul", 0)])
        self.assertEqual(blocks[1].items, ["Текст **важен**"])
        self.assertEqual(blocks[2].items, ["един", "два"])

    def test_nothing_is_executed_or_kept_from_html(self):
        text = '<p onclick="x()">Здравей <img src=x onerror=alert(1)> <script>alert(1)</script> <a href="javascript:x">линк</a></p>'
        blocks = report_text.parse(text)
        joined = " ".join(" ".join(b.items) for b in blocks)
        for bad in ("<", ">", "onclick", "onerror", "href"):
            self.assertNotIn(bad, joined)
        self.assertIn("Здравей", joined)
        self.assertIn("линк", joined)

    def test_comparison_signs_in_text_survive(self):
        self.assertEqual(report_text.parse("орбис < 3° и > 1°")[0].items, ["орбис < 3° и > 1°"])

    def test_inline_runs(self):
        self.assertEqual(report_text.inline_runs("a **b** c *d* snake_case _e_"),
                         [("a ", False, False), ("b", True, False), (" c ", False, False), ("d", False, True),
                          (" snake_case ", False, False), ("e", False, True)])
        self.assertEqual(report_text.plain("**Удебелено** и *курсив*"), "Удебелено и курсив")

    def test_markdown_is_clean_and_stable(self):
        once = report_text.to_markdown(MIXED)
        self.assertNotIn("<", once)
        self.assertNotIn("━", once)
        self.assertIn("## Подзаглавие", once)
        self.assertIn("- първа точка", once)
        self.assertIn("1. едно", once)
        self.assertEqual(report_text.to_markdown(once), once)          # втори прочит не променя нищо

    def test_period_sections(self):
        sections = report_text.split_sections(MIXED)
        self.assertEqual([s[0] for s in sections], ["Общ преглед на периода", "Октомври 2026"])
        self.assertEqual(report_text.split_sections("Само текст")[0], ("", "Само текст"))
        self.assertEqual(report_text.split_sections("Увод\n<h2>Месец</h2>\nТекст")[0], ("", "Увод"))


class ExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(main.app)

    def setUp(self):
        limiter.reset()

    def report(self, email, body=None, text="<p>текст</p>"):
        headers = register_and_login(self.client, email)
        fake = mock.AsyncMock(return_value=text)
        with mock.patch.object(main.ai_interpreter, "interpret_chart", fake):
            job = self.client.post("/jobs?wait=60", json=body or CHART, headers=headers).json()["job"]
        self.assertEqual(job["status"], "succeeded", job)
        return headers, job["report_id"]

    def export(self, headers, report_id, fmt="docx"):
        return self.client.get(f"/reports/{report_id}/export", params={"format": fmt}, headers=headers)

    @staticmethod
    def docx_text(response):
        doc = Document(io.BytesIO(response.content))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts += [cell.text for cell in row.cells]
        return "\n".join(parts)

    def test_docx_has_the_text_the_summary_and_no_raw_markup(self):
        headers, rid = self.report("ex-docx@test.bg", {**CHART, "name": "Иван Иванов"}, MIXED)
        r = self.export(headers, rid)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["content-type"], "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(r.content)))
        text = self.docx_text(r)
        for expected in ("Астрологичен доклад", "Иван Иванов", "Планетарни позиции", "Домове", "Асцендент", "Аспекти",
                         "Подзаглавие", "първа точка", "Първи абзац с удебелено и курсив."):
            self.assertIn(expected, text)
        for raw in ("**", "<h2>", "<p>", "<strong>", "━"):
            self.assertNotIn(raw, text)

    def test_docx_is_a4_with_the_period_sections_in_the_contents(self):
        headers, rid = self.report("ex-a4@test.bg", None, MIXED)
        doc = Document(io.BytesIO(self.export(headers, rid).content))
        section = doc.sections[0]
        self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
        self.assertAlmostEqual(section.page_height.cm, 29.7, places=1)
        text = self.docx_text(self.export(headers, rid))
        self.assertIn("Съдържание", text)
        self.assertIn("1. Общ преглед на периода", text)
        self.assertIn("2. Октомври 2026", text)

    def test_docx_without_a_birth_time_has_no_houses_or_ascendant(self):
        headers, rid = self.report("ex-notime@test.bg", {"date": "1990-05-15", "lat": 42.6977, "lon": 23.3219,
                                                          "birth_time_known": False}, "<p>текст</p>")
        text = self.docx_text(self.export(headers, rid))
        self.assertNotIn("\nАсцендент\n", text)                      # няма ред „Асцендент“ в таблицата (обяснението го споменава)
        self.assertNotIn("празен", text)
        self.assertNotIn("Домове\n", text)
        self.assertIn("часът е неизвестен", text)
        self.assertIn("Козирог или Водолей (според часа)", text)

    def test_two_people_get_two_charts_and_the_relationship(self):
        body = {**CHART, "name": "Иван", "partner_name": "Мария", "partner_date": "1992-07-20", "partner_time": "09:30",
                "partner_lat": 43.2141, "partner_lon": 27.9147, "relationship": "friend"}
        headers, rid = self.report("ex-pair@test.bg", body)
        text = self.docx_text(self.export(headers, rid))
        for expected in ("Карта: Иван", "Карта: Мария", "Приятелство"):
            self.assertIn(expected, text)

    def test_markdown_export(self):
        headers, rid = self.report("ex-md@test.bg", {**CHART, "name": "Иван"}, MIXED)
        r = self.export(headers, rid, "md")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.headers["content-type"].startswith("text/markdown"))
        text = r.content.decode("utf-8")
        self.assertTrue(text.startswith("# Общ анализ"), text[:80])
        for expected in ("**Иван**", "## Октомври 2026", "- първа точка", "1. едно"):
            self.assertIn(expected, text)
        for raw in ("<h2>", "<p>", "━"):
            self.assertNotIn(raw, text)

    def test_the_file_name_has_no_personal_data(self):
        headers, rid = self.report("ex-name@test.bg", {**CHART, "name": "Иван Иванов"})
        for fmt in ("docx", "md"):
            disposition = self.export(headers, rid, fmt).headers["content-disposition"]
            self.assertRegex(disposition, rf'^attachment; filename="AstroMind-general-\d{{4}}-\d{{2}}-\d{{2}}-{rid}\.{fmt}"$')
            self.assertNotIn("Иван", disposition)
            self.assertNotIn("ex-name", disposition)

    def test_only_the_owner_can_export(self):
        headers, rid = self.report("ex-owner@test.bg")
        other = register_and_login(self.client, "ex-other@test.bg")
        self.assertEqual(self.export(other, rid).status_code, 404)
        self.assertEqual(self.client.get(f"/reports/{rid}/export").status_code, 401)
        self.assertEqual(self.export(headers, rid + 999).status_code, 404)
        self.assertEqual(self.export(headers, rid, "pdf").status_code, 400)

    def test_export_costs_nothing_and_never_calls_the_ai(self):
        headers, rid = self.report("ex-free@test.bg")
        before = self.client.get("/jobs", headers=headers).json()["jobs"]
        boom = mock.AsyncMock(side_effect=AssertionError("износът не бива да вика AI"))
        with mock.patch.object(main.ai_interpreter, "_call_api", boom), \
                mock.patch.object(main.ai_interpreter, "interpret_chart", boom):
            for fmt in ("docx", "md", "docx"):
                self.assertEqual(self.export(headers, rid, fmt).status_code, 200)
        self.assertEqual(self.client.get("/jobs", headers=headers).json()["jobs"], before)
        boom.assert_not_called()

    def test_an_old_report_without_input_data_still_exports(self):
        headers = register_and_login(self.client, "ex-old@test.bg")
        from database import SessionLocal, User
        db = SessionLocal()
        try:
            uid = db.query(User.id).filter(User.email == "ex-old@test.bg").scalar()
            r = data_api.save_report(db, db.get(User, uid), content="**Стар** отчет", report_type="love", profile_name="Стар",
                                     label="Любов", cost_cents=0, params={"imported": True})
            db.commit()
            rid = r.id
        finally:
            db.close()
        text = self.docx_text(self.export(headers, rid))
        self.assertIn("Стар отчет", text)
        self.assertNotIn("Планетарни позиции", text)
        self.assertIn("**Стар** отчет", self.export(headers, rid, "md").content.decode("utf-8"))

    def test_rate_limited(self):
        headers, rid = self.report("ex-rate@test.bg")
        with mock.patch.object(data_api, "EXPORT_LIMIT_PER_HOUR", 2):
            statuses = [self.export(headers, rid, "md").status_code for _ in range(3)]
        self.assertEqual(statuses, [200, 200, 429])


if __name__ == "__main__":
    unittest.main()
