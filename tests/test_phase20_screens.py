"""
Фаза 20: екраните от одита (AM-N06, AM-N07 и текстовете от AM-N09).

- баланс на елементите при равен брой (2/1/2/2 е баланс, не „преобладава огън“);
- таблица в текста на отчета: един разбор в Python (износ) и JS (екран), DOCX и Markdown;
- История: вид анализ, момент или период и вторият човек;
- натален и моментен отчет с една тема не са с еднакво име.
"""
import io
import json
import os
import shutil
import subprocess
import unittest
import zipfile
from datetime import datetime
from unittest import mock

import testenv  # noqa: F401  (трябва да е преди main)
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import data_api  # noqa: E402
import insights  # noqa: E402
import main  # noqa: E402
import report_export  # noqa: E402
import report_text  # noqa: E402
from database import Report  # noqa: E402
from docx_generator import DOCXGenerator  # noqa: E402
from rate_limit import limiter  # noqa: E402
from testenv import CHART, register_and_login  # noqa: E402

FRONTEND = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "src", "utils", "reportText.js")


def chart_with(signs):
    return {"planets": {name: {"zodiac_sign": sign} for name, sign in signs.items()}, "angles": {}}


class ElementBalanceTest(unittest.TestCase):
    def test_two_one_two_two_is_a_balance_between_the_leading_elements(self):
        chart = chart_with({"Sun": "Aries", "Moon": "Leo", "Mercury": "Gemini", "Venus": "Libra", "Mars": "Cancer",
                            "Jupiter": "Scorpio", "Saturn": "Taurus"})
        result = insights.big_three(chart, time_known=False)
        self.assertEqual({e["id"]: e["count"] for e in result["elements"]}, {"fire": 2, "earth": 1, "air": 2, "water": 2})
        self.assertTrue(result["balanced"])
        self.assertEqual([e["id"] for e in result["dominant_elements"]], ["fire", "air", "water"])
        self.assertIn("Няма един преобладаващ елемент", result["balance_text"])
        self.assertIn("огън, въздух и вода", result["balance_text"])
        self.assertIn("7 тела", result["method"])

    def test_a_single_leader_is_still_the_dominant_element(self):
        chart = chart_with({"Sun": "Aries", "Moon": "Leo", "Mercury": "Sagittarius", "Venus": "Libra", "Mars": "Cancer",
                            "Jupiter": "Scorpio", "Saturn": "Taurus"})
        result = insights.big_three(chart, time_known=False)
        self.assertFalse(result["balanced"])
        self.assertEqual(result["dominant_element"]["id"], "fire")
        self.assertEqual(result["balance_text"], "")


TABLE = "| Дата | Събитие |\n| --- | --- |\n| 10 октомври | Новолуние във Везни |\n| 24 октомври | **Марс** в опозиция |"


class TableParseTest(unittest.TestCase):
    def test_consecutive_pipe_lines_are_one_table_with_a_header(self):
        blocks = report_text.parse("Преди.\n\n" + TABLE + "\n\nСлед.")
        self.assertEqual([b.kind for b in blocks], ["p", "table", "p"])
        table = blocks[1]
        self.assertTrue(table.header)
        self.assertEqual(table.rows, [["Дата", "Събитие"], ["10 октомври", "Новолуние във Везни"], ["24 октомври", "**Марс** в опозиция"]])

    def test_a_table_without_a_separator_row_has_no_header_and_text_after_it_is_a_new_paragraph(self):
        blocks = report_text.parse("| а | б |\n| в | г |\nТекст след таблицата.")
        self.assertEqual([b.kind for b in blocks], ["table", "p"])
        self.assertFalse(blocks[0].header)
        self.assertEqual(blocks[1].items, ["Текст след таблицата."])

    def test_a_single_pipe_line_stays_a_paragraph(self):
        blocks = report_text.parse("| само един ред |")
        self.assertEqual([(b.kind, b.items) for b in blocks], [("p", ["| само един ред |"])])

    def test_the_table_is_exported_to_markdown_with_a_separator(self):
        markdown = report_text.to_markdown(TABLE)
        self.assertEqual(markdown.splitlines()[:3], ["| Дата | Събитие |", "| --- | --- |", "| 10 октомври | Новолуние във Везни |"])

    def test_html_in_a_cell_is_not_kept(self):
        blocks = report_text.parse("| а | б |\n| <script>x()</script>текст | <b onclick=y()>жирно</b> |")
        self.assertEqual(blocks[0].rows[1], ["x()текст", "**жирно**"])

    def test_docx_has_a_real_table_with_a_bold_header(self):
        report = Report(user_id=1, profile_name="Тест", report_type="general", label="Таблица", status="completed",
                        content="<h2>Общ преглед</h2>\n" + TABLE, cost_cents=0, params={"name": "Тест", "date": "1990-02-15",
                                                                                         "time": "13:00", "lat": 42.7, "lon": 23.3})
        report.created_at = datetime(2026, 10, 10, 20, 0)
        doc = Document(io.BytesIO(report_export.docx_bytes(report)))
        tables = [t for t in doc.tables if t.rows and t.rows[0].cells[0].text == "Дата"]
        self.assertEqual(len(tables), 1)
        table = tables[0]
        self.assertEqual([[c.text for c in row.cells] for row in table.rows][1], ["10 октомври", "Новолуние във Везни"])
        self.assertTrue(all(run.bold for run in table.rows[0].cells[0].paragraphs[0].runs))
        self.assertEqual(table.rows[2].cells[1].text, "Марс в опозиция")


@unittest.skipUnless(shutil.which("node") and os.path.exists(FRONTEND), "Node.js или файлът на екрана не е наличен")
class ScreenAndExportParseTheSameTest(unittest.TestCase):
    """Екранът (src/utils/reportText.js) и износът (report_text.py) четат текста еднакво."""

    SAMPLES = [
        "# Заглавие\n\nАбзац **удебелен** и *курсив*.\n\n- а\n- б\n\n1. едно\n2. две\n\n---\n\nКрай",
        "<h2>Общ преглед</h2>\n" + TABLE + "\n\nСлед таблицата.",
        "| а | б |\n| в | г |\nТекст след таблицата.",
        "| само един ред |",
        "Текст\n| x | y |\n|---|---|\n| 1 | 2 |\nОтново текст",
        "<p>Здравей <b>свят</b></p><ul><li>един</li><li>два</li></ul>",
        "",
    ]

    def js_blocks(self):
        script = ("import { parseReport } from %s;\nconst samples = JSON.parse(process.argv[1]);\n"
                  "console.log(JSON.stringify(samples.map((s) => parseReport(s))));") % json.dumps("file://" + FRONTEND)
        out = subprocess.run(["node", "--input-type=module", "-e", script, json.dumps(self.SAMPLES)],
                             capture_output=True, text=True, timeout=60)
        self.assertEqual(out.returncode, 0, out.stderr)
        return json.loads(out.stdout)

    def test_the_same_blocks_for_every_sample(self):
        for sample, js in zip(self.SAMPLES, self.js_blocks()):
            python = [{"kind": b.kind, "level": b.level, "items": b.items, "rows": b.rows, "header": b.header}
                      for b in report_text.parse(sample)]
            normalized = [{"kind": b["kind"], "level": b["level"], "items": b["items"], "rows": b.get("rows", []),
                           "header": bool(b.get("header", False))} for b in js]
            self.assertEqual(python, normalized, sample)


class HistoryLineTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(main.app)

    def setUp(self):
        limiter.reset()

    def run_job(self, headers, body):
        limiter.reset()                                  # в тестовете лимитът е два анализа на час
        with mock.patch.object(main.ai_interpreter, "interpret_chart", mock.AsyncMock(return_value="<p>текст</p>")):
            return self.client.post("/jobs?wait=60", json=body, headers=headers).json()["job"]

    def test_natal_snapshot_and_partner_reports_differ_in_label_and_summary(self):
        h = register_and_login(self.client, "p20-history@test.bg")
        natal = self.run_job(h, {**CHART, "name": "Аз", "report_type": "career"})
        snapshot = self.run_job(h, {**CHART, "name": "Аз", "report_type": "career", "target_date": "2026-10-10", "target_time": "12:00"})
        pair = self.run_job(h, {**CHART, "name": "Аз", "report_type": "career", "partner_name": "Иван", "partner_date": "1985-07-12",
                                "partner_time": "18:45", "partner_lat": 48.2082, "partner_lon": 16.3738})
        rows = {r["id"]: r for r in self.client.get("/reports", headers=h).json()}
        n, s, p = rows[natal["report_id"]], rows[snapshot["report_id"]], rows[pair["report_id"]]
        self.assertNotEqual(n["label"], s["label"])
        self.assertTrue(s["label"].endswith("за дата"), s["label"])
        self.assertEqual((n["mode"], n["when"], n["partner"]), ("natal", "", ""))
        self.assertEqual((s["mode"], s["when"]), ("snapshot", "10.10.2026 12:00"))
        self.assertEqual((p["mode"], p["partner"]), ("natal", "Иван"))
        self.assertIn("Иван", p["label"])

    def test_the_summary_of_a_period_report_names_the_dates(self):
        info = data_api.report_mode({"is_dynamic": True, "target_date": "2026-10-01", "end_date": "2026-11-30"})
        self.assertEqual(info, {"mode": "period", "when": "01.10.2026 – 30.11.2026", "partner": ""})
        self.assertEqual(data_api.report_mode(None), {"mode": "natal", "when": "", "partner": ""})


if __name__ == "__main__":
    unittest.main()
