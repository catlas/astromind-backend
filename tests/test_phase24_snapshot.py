"""
Фаза 24 (регресионен одит 11.10.2026, R06): неизменяем отчет и пълни изчисления.

- при успешно генериране се записва изчислителен снимка (карти, небе/календар, UTC, зони, координати, версии, отпечатък на текста);
- износът (DOCX, Markdown) е от снимката: двигателят и скенерът не се викат, а промяна на входните данни не променя износа;
- target_lat/target_lon и профилите се записват; рождената и прогнозната зона са различни в износа;
- липсващо тяло (Хирон) отказва пълния анализ преди AI, такса и запис; датата извън файловете с ефемериди се отказва;
- над полярния кръг домовете са Porphyry и това е видимо; старите отчети са „legacy“ без измислен снимка;
- миграция 0014 е обратима.
"""
import io
import os
import unittest
import zipfile
from unittest import mock

import testenv  # noqa: F401  (трябва да е преди main)
import sqlalchemy as sa  # noqa: E402
from alembic import command  # noqa: E402
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import engine  # noqa: E402
import generation  # noqa: E402
import main  # noqa: E402
import report_export  # noqa: E402
import report_snapshot  # noqa: E402
import scanner  # noqa: E402
from database import Report, SessionLocal  # noqa: E402
from rate_limit import limiter  # noqa: E402
from schemas import ChartRequest  # noqa: E402
from test_euro_balance_migration import alembic_config, make_engine  # noqa: E402
from testenv import CHART, OVERVIEW_OK, register_and_login  # noqa: E402

VIENNA = {"date": "1985-07-12", "time": "18:45", "lat": 48.2082, "lon": 16.3738}
NEW_YORK = {"target_lat": 40.7128, "target_lon": -74.0060}
SNAPSHOT_BODY = {**VIENNA, "name": "Иван", "target_date": "2026-10-10", "target_time": "12:00", "forecast_timezone": "Europe/Sofia", **NEW_YORK}
PERIOD_BODY = {**VIENNA, "name": "Иван", "is_dynamic": True, "target_date": "2026-10-11", "end_date": "2026-10-31",
               "forecast_timezone": "Europe/Sofia"}


def latest_report(email):
    db = SessionLocal()
    try:
        from database import User
        uid = db.query(User.id).filter(User.email == email).scalar()
        row = db.query(Report).filter(Report.user_id == uid).order_by(Report.id.desc()).first()
        return row.id if row else None
    finally:
        db.close()


def load(report_id):
    db = SessionLocal()
    try:
        row = db.get(Report, report_id)
        db.expunge(row)
        return row
    finally:
        db.close()


def docx_text(body: bytes) -> str:
    doc = Document(io.BytesIO(body))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts += [cell.text for cell in row.cells]
    return "\n".join(parts)


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(main.app)

    def setUp(self):
        limiter.reset()

    def generate(self, email, body):
        headers = register_and_login(self.client, email)
        month = mock.patch.object(main.ai_interpreter, "_process_monthly_chunk", mock.AsyncMock(return_value="<p>Месец</p>"))
        overview = mock.patch.object(main.ai_interpreter, "compose_period_overview", mock.AsyncMock(return_value=OVERVIEW_OK))
        single = mock.patch.object(main.ai_interpreter, "interpret_chart", mock.AsyncMock(return_value="<p>Анализ на момента.</p>"))
        with month, overview, single:
            response = self.client.post("/jobs?wait=60", json=body, headers=headers)
        self.assertEqual(response.status_code, 200, response.text)
        job = response.json()["job"]
        self.assertEqual(job["status"], "succeeded", job)
        return headers, job["report_id"]


class SnapshotSavedTest(Base):
    def test_a_natal_report_keeps_the_calculation_it_was_written_from(self):
        _, report_id = self.generate("p24-natal@test.bg", {**VIENNA, "name": "Иван"})
        snap = load(report_id).snapshot
        self.assertEqual((snap["version"], snap["mode"]), (report_snapshot.SNAPSHOT_VERSION, "natal"))
        person = report_snapshot.person(snap, "user")
        self.assertEqual((person["birth"]["lat"], person["birth"]["lon"], person["birth"]["zone"]), (48.2082, 16.3738, "Europe/Vienna"))
        self.assertEqual(person["birth"]["utc"], person["chart"]["datetime_utc"])
        fresh = engine.get_engine().calculate_chart(**VIENNA)
        self.assertEqual(person["chart"]["planets"]["Sun"]["longitude"], fresh["planets"]["Sun"]["longitude"])
        self.assertTrue(person["natal_aspects"])
        self.assertEqual(snap["ephemeris"]["missing"], [])
        self.assertEqual(snap["ephemeris"]["years"], [1800, 2399])
        self.assertEqual(set(snap["versions"]), {"snapshot", "calendar", "factpack", "orbs", "checker", "prompt"})

    def test_the_text_fingerprint_belongs_to_the_saved_text(self):
        _, report_id = self.generate("p24-hash@test.bg", {**VIENNA, "name": "Иван"})
        row = load(report_id)
        self.assertTrue(report_snapshot.verify(row.snapshot, row.content))
        self.assertFalse(report_snapshot.verify(row.snapshot, row.content + " друг текст"))

    def test_a_relocated_moment_keeps_its_own_coordinates_zone_and_utc(self):
        _, report_id = self.generate("p24-reloc@test.bg", SNAPSHOT_BODY)
        row = load(report_id)
        snap = row.snapshot
        target = snap["target"]
        self.assertEqual((snap["mode"], snap["forecast_zone"], target["relocated"]), ("snapshot", "Europe/Sofia", True))
        self.assertEqual((target["lat"], target["lon"]), (40.7128, -74.006))
        self.assertEqual(target["chart"]["location"], {"latitude": 40.7128, "longitude": -74.006})
        self.assertEqual(target["utc"], target["chart"]["datetime_utc"])
        self.assertEqual((row.params["target_lat"], row.params["target_lon"]), (40.7128, -74.006))
        self.assertEqual(row.params["forecast_timezone"], "Europe/Sofia")
        self.assertTrue(target["aspects"]["user"])

    def test_a_period_report_keeps_the_calendar_and_the_zone(self):
        _, report_id = self.generate("p24-period@test.bg", PERIOD_BODY)
        snap = load(report_id).snapshot
        self.assertEqual(snap["mode"], "period")
        self.assertEqual((snap["period"]["start"], snap["period"]["end"], snap["period"]["zone"]), ("2026-10-11", "2026-10-31", "Europe/Sofia"))
        self.assertEqual(snap["calendar"]["version"], scanner.CALENDAR_VERSION)
        self.assertTrue(snap["calendar"]["events"])
        self.assertFalse([k for e in snap["calendar"]["events"] for k in e if k.startswith("_")])        # без вътрешните полета

    def test_two_people_and_the_relationship_are_kept(self):
        body = {**VIENNA, "name": "Иван", "partner_name": "Мария", "partner_date": "1992-07-20", "partner_time": "09:30",
                "partner_lat": 43.2141, "partner_lon": 27.9147, "relationship": "friend"}
        _, report_id = self.generate("p24-pair@test.bg", body)
        snap = load(report_id).snapshot
        self.assertEqual([p["role"] for p in snap["people"]], ["user", "partner"])
        self.assertEqual(report_snapshot.person(snap, "partner")["birth"]["zone"], "Europe/Sofia")
        self.assertEqual(snap["relationship"], {"kind": "friend", "parent": None})


class ExportFromSnapshotTest(Base):
    def test_the_export_does_not_use_the_current_engine_or_the_current_input(self):
        headers, report_id = self.generate("p24-export@test.bg", SNAPSHOT_BODY)
        before_md = self.client.get(f"/reports/{report_id}/export?format=md", headers=headers).text
        before_docx = docx_text(self.client.get(f"/reports/{report_id}/export?format=docx", headers=headers).content)
        # променен двигател и редактирани входни данни на отчета: износът остава същият
        db = SessionLocal()
        try:
            row = db.get(Report, report_id)
            row.params = {**row.params, "date": "1999-01-01", "lat": 0.0, "lon": 0.0, "target_lat": 1.0}
            db.commit()
        finally:
            db.close()
        with mock.patch.object(engine.AstrologyEngine, "calculate_chart", side_effect=AssertionError("двигателят не се вика")), \
                mock.patch.object(scanner.TransitScanner, "build_calendar", side_effect=AssertionError("скенерът не се вика")):
            after_md = self.client.get(f"/reports/{report_id}/export?format=md", headers=headers).text
            after_docx = docx_text(self.client.get(f"/reports/{report_id}/export?format=docx", headers=headers).content)
        self.assertEqual(after_md, before_md)
        self.assertEqual(after_docx, before_docx)

    def test_birth_zone_forecast_zone_and_moment_are_separate_in_every_format(self):
        headers, report_id = self.generate("p24-zones@test.bg", SNAPSHOT_BODY)
        md = self.client.get(f"/reports/{report_id}/export?format=md", headers=headers).text
        self.assertIn("Рождена зона: Иван: Europe/Vienna (UTC+2)", md)
        self.assertIn("Прогнозна зона: Europe/Sofia", md)
        self.assertRegex(md, r"Момент на прогнозата: 10\.10\.2026 12:00 \(UTC\+3, Europe/Sofia\) · UTC 2026-10-10 09:00")
        self.assertIn("небе за 40.7128, -74.006 (релокация)", md)            # небето е за Ню Йорк, не за рождените координати
        docx = docx_text(self.client.get(f"/reports/{report_id}/export?format=docx", headers=headers).content)
        for needle in ("Рождена зона", "Europe/Vienna", "Прогнозна зона", "Europe/Sofia", "Момент на прогнозата"):
            self.assertIn(needle, docx)

    def test_markdown_has_the_context_and_the_calculation_tables(self):
        headers, report_id = self.generate("p24-tables@test.bg", SNAPSHOT_BODY)
        md = self.client.get(f"/reports/{report_id}/export?format=md", headers=headers).text
        for needle in ("## Рождени данни", "| Човек | Дата и час | Координати | Рождена зона | UTC момент |", "| Тяло | Позиция | Дом | Движение |",
                       "| Двойка | Аспект | Орбис |", "### Небе на 10.10.2026 12:00 (UTC+3, транзити)", "48.2082, 16.3738",
                       "Ефемериди: swiss", "Отпечатък на текста: sha256"):
            self.assertIn(needle, md, msg=needle)
        self.assertNotIn("legacy", md)

    def test_a_period_export_has_the_calendar_table(self):
        headers, report_id = self.generate("p24-cal@test.bg", PERIOD_BODY)
        md = self.client.get(f"/reports/{report_id}/export?format=md", headers=headers).text
        self.assertIn("### Календар на периода", md)
        self.assertIn("| Дата и час | Вид | Събитие |", md)
        self.assertRegex(md, r"\| \d{2}\.\d{2}\.2026 \d{2}:\d{2} \| (Станция|Влизане в знак|Лунация|Аспект) \|")
        docx = docx_text(self.client.get(f"/reports/{report_id}/export?format=docx", headers=headers).content)
        self.assertIn("Календар на периода", docx)

    def test_the_docx_is_a_valid_zip_and_keeps_the_notes(self):
        headers, report_id = self.generate("p24-zip@test.bg", SNAPSHOT_BODY)
        body = self.client.get(f"/reports/{report_id}/export?format=docx", headers=headers).content
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(body)))
        self.assertIn("Снимка на изчисленията", docx_text(body))


class LegacyTest(Base):
    def make_legacy(self, email):
        headers = register_and_login(self.client, email)
        db = SessionLocal()
        try:
            from database import User
            uid = db.query(User.id).filter(User.email == email).scalar()
            row = Report(user_id=uid, report_type="general", label="Общ анализ", content="Стар текст.", status="completed",
                         params={"name": "Иван", **VIENNA})
            db.add(row)
            db.commit()
            return headers, row.id
        finally:
            db.close()

    def test_an_old_report_has_no_invented_snapshot_and_says_so(self):
        headers, report_id = self.make_legacy("p24-legacy@test.bg")
        row = load(report_id)
        self.assertIsNone(row.snapshot)
        self.assertTrue(report_snapshot.is_legacy(row.snapshot))
        md = self.client.get(f"/reports/{report_id}/export?format=md", headers=headers).text
        self.assertIn("Стар отчет (legacy)", md)
        self.assertNotIn("Снимка на изчисленията", md)
        self.assertIn("legacy", docx_text(self.client.get(f"/reports/{report_id}/export?format=docx", headers=headers).content))

    def test_the_api_marks_old_and_new_reports(self):
        headers, legacy_id = self.make_legacy("p24-mark@test.bg")
        listing = {r["id"]: r for r in self.client.get("/reports", headers=headers).json()}
        self.assertTrue(listing[legacy_id]["legacy"])
        self.assertTrue(self.client.get(f"/reports/{legacy_id}", headers=headers).json()["legacy"])


class MandatoryBodiesTest(Base):
    def test_a_missing_body_refuses_the_analysis_before_the_ai(self):
        with self.assertRaises(generation.GenerationFailure) as caught:
            generation.require_complete_charts({"ephemeris": {"missing": ["Chiron"]}})
        self.assertEqual(caught.exception.code, "ephemeris_incomplete")
        self.assertIn("Хирон", caught.exception.message)
        generation.require_complete_charts({"ephemeris": {"missing": []}}, None, {"ephemeris": {}})            # пълните карти минават

    def test_a_real_chart_without_chiron_is_refused_and_nothing_is_taken(self):
        headers = register_and_login(self.client, "p24-chiron@test.bg")
        original = engine.AstrologyEngine._planet_positions

        def without_chiron(self_, jd):
            planets = original(self_, jd)
            planets["Chiron"]["longitude"] = None
            return planets

        ai = mock.AsyncMock(return_value="<p>текст</p>")
        with mock.patch.object(engine.AstrologyEngine, "_planet_positions", without_chiron), \
                mock.patch.object(main.ai_interpreter, "interpret_chart", ai):
            response = self.client.post("/jobs?wait=60", json={**VIENNA, "name": "Иван"}, headers=headers)
        self.assertEqual(response.status_code, 503, response.text)
        self.assertIn("Хирон", response.json()["detail"])
        self.assertEqual(ai.await_count, 0)
        self.assertEqual(self.client.get("/jobs", headers=headers).json()["jobs"], [])
        self.assertIsNone(latest_report("p24-chiron@test.bg"))

    def test_an_unknown_birth_time_is_not_a_missing_body(self):
        chart = engine.get_engine().calculate_chart("1985-07-12", "12:00", 48.2082, 16.3738, time_known=False)
        generation.require_complete_charts(chart)


class SupportedRangeTest(Base):
    def test_the_range_comes_from_the_files_that_are_there(self):
        self.assertEqual(engine.supported_year_range(), (1800, 2399))
        self.assertEqual(len(engine.ephemeris_fingerprint()), 16)

    def test_a_date_outside_the_files_is_refused_with_a_clear_message(self):
        headers = register_and_login(self.client, "p24-range@test.bg")
        for field, value in (("date", "1700-01-01"), ("target_date", "2500-01-01")):
            response = self.client.post("/jobs?wait=60", json={**CHART, field: value, **({"target_time": "12:00"} if field == "target_date" else {})},
                                        headers=headers)
            self.assertEqual(response.status_code, 400, response.text)
            self.assertIn("1800–2399", response.json()["detail"])


class HouseSystemTest(Base):
    def test_a_polar_birth_place_uses_porphyry_and_says_so(self):
        tromso = engine.get_engine().calculate_chart("1990-02-15", "13:00", 69.65, 18.96)
        sofia = engine.get_engine().calculate_chart("1990-02-15", "13:00", 42.6977, 23.3219)
        self.assertEqual((tromso["house_system"], sofia["house_system"]), ("porphyry_fallback", "placidus"))
        self.assertEqual(len(tromso["houses"]), 12)
        self.assertIn("Porphyry", generation.house_system_note([("Иван", tromso)]))
        self.assertEqual(generation.house_system_note([("Иван", sofia)]), "")

    def test_the_note_is_in_the_report_and_in_the_export(self):
        body = {"date": "1990-02-15", "time": "13:00", "lat": 69.65, "lon": 18.96, "name": "Иван"}
        headers, report_id = self.generate("p24-polar@test.bg", body)
        row = load(report_id)
        self.assertTrue(row.content.startswith("**Бележка за домовете.**"))
        self.assertEqual(row.snapshot["house_systems"], {"user": "porphyry_fallback"})
        md = self.client.get(f"/reports/{report_id}/export?format=md", headers=headers).text
        self.assertIn("Porphyry", md)


class MigrationTest(unittest.TestCase):
    def test_the_snapshot_column_is_added_empty_and_can_be_removed(self):
        db_engine = make_engine("p24")
        with db_engine.begin() as conn:
            command.upgrade(alembic_config(conn), "0013_memory_profile_id")
            conn.execute(sa.text("INSERT INTO users (id, email, hashed_password, full_name) VALUES (1, 'a@mig.bg', 'x', 'Тест')"))
            conn.execute(sa.text("INSERT INTO reports (id, user_id, report_type, label, content, cost_cents, status) "
                                 "VALUES (1, 1, 'general', 'Стар', 'текст', 0, 'completed')"))
            command.upgrade(alembic_config(conn), "0014_report_snapshot")
            self.assertIn("snapshot", {c["name"] for c in sa.inspect(conn).get_columns("reports")})
            self.assertIsNone(conn.execute(sa.text("SELECT snapshot FROM reports WHERE id = 1")).scalar())      # старият отчет остава legacy
            command.downgrade(alembic_config(conn), "0013_memory_profile_id")
            self.assertNotIn("snapshot", {c["name"] for c in sa.inspect(conn).get_columns("reports")})
            self.assertEqual(conn.execute(sa.text("SELECT content FROM reports WHERE id = 1")).scalar(), "текст")


if __name__ == "__main__":
    unittest.main()
