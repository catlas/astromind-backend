"""
Тестове за Фаза 12: неизвестен час на раждане.

Часът не се измисля: без час няма Асцендент, МС, домове и управители на домове, Луната няма позиция и аспекти, а телата,
които сменят знак през деня, се описват с възможните знаци. Пускане: python -m unittest discover -s tests -p "test_unknown_time.py"
"""
import unittest
from datetime import datetime, timedelta
from unittest import mock
from zoneinfo import ZoneInfo

import swisseph as swe  # type: ignore

import testenv  # noqa: F401  (трябва да е преди main)
from fastapi.testclient import TestClient  # noqa: E402
from pydantic import ValidationError  # noqa: E402

import aspects_engine  # noqa: E402
import birthtime  # noqa: E402
import engine  # noqa: E402
import factpack  # noqa: E402
import main  # noqa: E402
import text_check  # noqa: E402
from ai_interpreter import AIInterpreter  # noqa: E402
from rate_limit import limiter  # noqa: E402
from schemas import ChartRequest  # noqa: E402
from testenv import CHART, register_and_login  # noqa: E402

SOFIA = (42.6977, 23.3219)
SIGNS = ["Aries", "Taurus", "Gemini", "Cancer", "Leo", "Virgo", "Libra", "Scorpio", "Sagittarius", "Capricorn",
         "Aquarius", "Pisces"]


def independent_sign(local_midnight: datetime, body: int) -> str:
    """Знакът на тяло в даден местен момент в София, изчислен директно със swisseph и zoneinfo (без кода на двигателя)."""
    utc = local_midnight.replace(tzinfo=ZoneInfo("Europe/Sofia")).astimezone(ZoneInfo("UTC"))
    jd = swe.julday(utc.year, utc.month, utc.day, utc.hour + utc.minute / 60.0)
    longitude = swe.calc_ut(jd, body, swe.FLG_SWIEPH)[0][0]
    return SIGNS[int(longitude // 30)]


class EngineTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.eng = engine.get_engine()

    def unknown(self, date, **kw):
        return self.eng.calculate_chart(date, "ignored", SOFIA[0], SOFIA[1], time_known=False, **kw)

    def test_no_houses_angles_or_house_numbers(self):
        chart = self.unknown("1990-05-15")
        self.assertIs(chart["time_known"], False)
        self.assertEqual(chart["houses"], {})
        self.assertEqual(chart["angles"], {})
        self.assertTrue(all(p["house"] is None for p in chart["planets"].values()))
        self.assertIn("час неизвестен", chart["datetime_local"])
        self.assertEqual(chart["timezone"], "Europe/Sofia")

    def test_a_known_time_chart_is_marked_known_and_unchanged(self):
        chart = self.eng.calculate_chart("1990-05-15", "14:30", SOFIA[0], SOFIA[1])
        self.assertIs(chart["time_known"], True)
        self.assertEqual(len(chart["houses"]), 12)
        self.assertTrue(chart["angles"]["Ascendant_formatted"])
        self.assertNotIn("sign_ranges", chart)

    def test_the_moon_has_no_position(self):
        moon = self.unknown("1990-05-15")["planets"]["Moon"]
        self.assertTrue(all(moon[k] is None for k in ("longitude", "speed", "distance", "zodiac_sign", "formatted_pos")))

    def test_the_other_bodies_are_the_local_noon_positions(self):
        noon = self.eng.calculate_chart("1990-05-15", "12:00", SOFIA[0], SOFIA[1])
        chart = self.unknown("1990-05-15")
        for name, planet in noon["planets"].items():
            if name != "Moon":
                with self.subTest(body=name):
                    self.assertEqual(chart["planets"][name]["longitude"], planet["longitude"])

    def test_the_time_argument_is_ignored(self):
        first = self.eng.calculate_chart("1990-05-15", "03:00", SOFIA[0], SOFIA[1], time_known=False)
        second = self.eng.calculate_chart("1990-05-15", "99:99", SOFIA[0], SOFIA[1], time_known=False)
        self.assertEqual(first, second)

    def test_moon_signs_match_an_independent_calculation(self):
        for date in ("1990-05-15", "1985-07-12", "2000-01-01", "2026-03-20", "1975-12-31", "2010-02-28"):
            with self.subTest(date=date):
                start = datetime.fromisoformat(date)
                first = independent_sign(start, swe.MOON)
                last = independent_sign(start + timedelta(days=1), swe.MOON)
                expected = [first] + ([last] if last != first else [])
                self.assertEqual(self.unknown(date)["sign_ranges"]["Moon"]["signs"], expected)

    def test_a_planet_that_changes_sign_during_the_day_is_listed(self):
        ranges = self.unknown("2026-03-20")["sign_ranges"]          # пролетното равноденствие: Слънцето влиза в Овен
        self.assertEqual(ranges["Sun"]["signs"], ["Pisces", "Aries"])
        self.assertIn("Pisces", ranges["Sun"]["from"])
        self.assertIn("Aries", ranges["Sun"]["to"])

    def test_a_day_without_a_sign_change_lists_only_the_moon(self):
        ranges = self.unknown("1985-07-12")["sign_ranges"]
        self.assertEqual(list(ranges), ["Moon"])
        self.assertEqual(ranges["Moon"]["signs"], ["Taurus"])

    def test_works_for_the_whole_range_of_dates_and_places(self):
        for date, place in (("1900-01-01", SOFIA), ("2026-10-25", SOFIA), ("1999-12-31", (40.7128, -74.006)),
                            ("2026-03-29", (-33.8688, 151.2093)), ("1975-07-01", (64.1466, -21.9426))):
            with self.subTest(date=date, place=place):
                chart = self.eng.calculate_chart(date, "00:00", place[0], place[1], time_known=False)
                self.assertIn("Moon", chart["sign_ranges"])
                self.assertEqual(chart["houses"], {})


class AspectsTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        eng = engine.get_engine()
        cls.unknown = eng.calculate_chart("1990-05-15", "", SOFIA[0], SOFIA[1], time_known=False)
        cls.known = eng.calculate_chart("1992-07-20", "09:30", 43.2141, 27.9147)
        cls.transit = eng.calculate_chart("2026-10-10", "12:00", SOFIA[0], SOFIA[1])

    def test_natal_aspects_have_no_moon_and_no_angles(self):
        points = {p for a in aspects_engine.calculate_natal_aspects(self.unknown) for p in (a["planet1"], a["planet2"])}
        self.assertTrue(points)
        self.assertFalse(points & {"Moon", "ASC", "MC"})

    def test_synastry_with_one_unknown_person_has_no_moon_or_angle_of_that_person(self):
        rows = aspects_engine.calculate_synastry_aspects(self.unknown, self.known)
        self.assertTrue(rows)
        self.assertNotIn("Moon", {r["planet1"] for r in rows})
        self.assertNotIn("ASC", {r["planet1"] for r in rows})          # Асцендентът на първия човек липсва
        reverse = aspects_engine.calculate_synastry_aspects(self.known, self.unknown)
        self.assertNotIn("Moon", {r["planet2"] for r in reverse})

    def test_transits_to_an_unknown_chart_skip_its_moon_and_angles(self):
        rows = aspects_engine.calculate_transit_aspects_to_natal(self.unknown, self.transit, include_angles=True)
        self.assertTrue(rows)
        self.assertFalse({r["natal_planet"] for r in rows} & {"Moon", "ASC", "MC"})


class FactPackTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        eng = engine.get_engine()
        cls.unknown = eng.calculate_chart("2026-03-20", "", SOFIA[0], SOFIA[1], time_known=False)
        cls.known = eng.calculate_chart("1992-07-20", "09:30", 43.2141, 27.9147)
        cls.transit = eng.calculate_chart("2026-10-10", "12:00", SOFIA[0], SOFIA[1])

    def test_the_view_has_no_houses_angles_or_house_numbers(self):
        view = factpack.natal_view(self.unknown)
        self.assertEqual(view["birth_time"], "unknown")
        self.assertNotIn("houses", view)
        self.assertNotIn("angles", view)
        self.assertTrue(all("house" not in p for p in view["planets"].values()))
        self.assertNotIn("Moon", view["planets"])

    def test_signs_that_depend_on_the_time_are_listed_and_have_no_degree(self):
        view = factpack.natal_view(self.unknown)
        self.assertEqual(view["time_dependent_signs"]["Sun"]["possible_signs"], ["Pisces", "Aries"])
        self.assertEqual(view["time_dependent_signs"]["Moon"]["possible_signs"], ["Aries"])
        self.assertNotIn("formatted_pos", view["planets"]["Sun"])
        self.assertIn("formatted_pos", view["planets"]["Mercury"])

    def test_the_known_view_is_unchanged(self):
        view = factpack.natal_view(self.known)
        self.assertEqual(len(view["houses"]), 12)
        self.assertIn("Ascendant_formatted", view["angles"])
        self.assertTrue(all(isinstance(p["house"], int) for p in view["planets"].values()))

    def test_has_houses(self):
        self.assertFalse(factpack.has_houses(self.unknown))
        self.assertTrue(factpack.has_houses(self.known))

    def test_the_natal_block_says_the_time_is_unknown(self):
        text = AIInterpreter._natal_block("Мария", self.unknown)
        self.assertIn("UNKNOWN", text)
        self.assertNotIn("Ascendant_formatted", text)
        self.assertNotIn('"houses"', text)
        known = AIInterpreter._natal_block("Иван", self.known)
        self.assertIn("Ascendant_formatted", known)
        self.assertNotIn("UNKNOWN", known)

    def test_overlays_exist_only_in_the_houses_of_a_person_with_a_known_time(self):
        both_unknown = AIInterpreter._overlay_blocks(self.unknown, self.unknown, "Мария", "Ана")
        self.assertEqual(both_unknown, "")
        text = AIInterpreter._overlay_blocks(self.unknown, self.known, "Мария", "Иван")
        self.assertNotIn("PARTNER PLANETS IN USER'S NATAL HOUSES", text)         # Мария няма домове
        self.assertIn("МАРИЯ PLANETS IN ИВАН'S NATAL HOUSES", text)                # Иван има
        reverse = AIInterpreter._overlay_blocks(self.known, self.unknown, "Иван", "Мария")
        self.assertIn("PARTNER PLANETS IN USER'S NATAL HOUSES", reverse)
        self.assertNotIn("ИВАН PLANETS IN МАРИЯ'S NATAL HOUSES", reverse)

    def test_transit_blocks_skip_the_houses_but_keep_the_aspects(self):
        text = AIInterpreter._transit_blocks("МАРИЯ", "Мария", self.unknown, self.transit)
        self.assertNotIn("NATAL HOUSES", text)
        self.assertIn("TRANSIT ASPECTS TO МАРИЯ'S NATAL CHART", text)
        with_time = AIInterpreter._transit_blocks("ИВАН", "Иван", self.known, self.transit)
        self.assertIn("NATAL HOUSES", with_time)


class RequestTest(unittest.TestCase):
    BASE = {"date": "1990-05-15", "lat": 42.6977, "lon": 23.3219}

    def test_a_known_time_is_required_unless_marked_unknown(self):
        with self.assertRaises(ValidationError):
            ChartRequest(**self.BASE)
        with self.assertRaises(ValidationError):
            ChartRequest(**self.BASE, time="  ")
        self.assertEqual(ChartRequest(**self.BASE, time="14:30").time, "14:30")

    def test_unknown_time_needs_no_time_and_clears_the_fold(self):
        request = ChartRequest(**self.BASE, birth_time_known=False, birth_fold=1, time="03:30")
        self.assertEqual(request.time, engine.UNKNOWN_TIME_ANCHOR)
        self.assertIsNone(request.birth_fold)

    def test_the_second_person_can_have_an_unknown_time(self):
        request = ChartRequest(**self.BASE, time="14:30", partner_date="1992-07-20", partner_lat=43.2, partner_lon=27.9,
                               partner_time_known=False, partner_fold=0)
        self.assertEqual(request.partner_time, engine.UNKNOWN_TIME_ANCHOR)
        self.assertIsNone(request.partner_fold)
        self.assertIs(request.birth_time_known, True)
        # без втори човек флагът не добавя час
        single = ChartRequest(**self.BASE, time="14:30", partner_time_known=False)
        self.assertIsNone(single.partner_time)

    def test_incomplete_data_for_the_second_person_is_refused_not_dropped(self):
        for partial in ({"partner_date": "1992-07-20"},
                        {"partner_date": "1992-07-20", "partner_time": "09:30"},
                        {"partner_date": "1992-07-20", "partner_lat": 43.2, "partner_lon": 27.9},
                        {"partner_time": "09:30", "partner_lat": 43.2, "partner_lon": 27.9}):
            with self.subTest(partial=partial):
                with self.assertRaises(ValidationError):
                    ChartRequest(**self.BASE, time="14:30", **partial)
        whole = ChartRequest(**self.BASE, time="14:30", partner_date="1992-07-20", partner_time="09:30",
                             partner_lat=43.2, partner_lon=27.9)
        self.assertEqual(whole.partner_time, "09:30")
        self.assertIsNone(ChartRequest(**self.BASE, time="14:30", partner_name="Само име").partner_date)


class NoteTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        eng = engine.get_engine()
        cls.unknown = eng.calculate_chart("1990-05-15", "", SOFIA[0], SOFIA[1], time_known=False)
        cls.equinox = eng.calculate_chart("2026-03-20", "", SOFIA[0], SOFIA[1], time_known=False)
        cls.known = eng.calculate_chart("1992-07-20", "09:30", 43.2141, 27.9147)

    def test_no_note_when_every_time_is_known(self):
        self.assertEqual(birthtime.note([("Иван", self.known)]), "")
        self.assertEqual(birthtime.block([]), "")

    def test_the_note_names_the_person_and_the_moon_signs(self):
        text = birthtime.note([("Мария", self.unknown)])
        self.assertIn("Мария", text)
        self.assertIn("Козирог или Водолей", text)
        self.assertIn("не е известен", text)
        self.assertIn("12:00", text)

    def test_a_certain_moon_sign_is_said_to_be_certain(self):
        eng = engine.get_engine()
        chart = eng.calculate_chart("1985-07-12", "", SOFIA[0], SOFIA[1], time_known=False)
        self.assertIn("Луната е в Телец (знакът е сигурен", birthtime.note([("Мария", chart)]))

    def test_a_planet_changing_sign_is_listed(self):
        text = birthtime.note([("Мария", self.equinox)])
        self.assertIn("Слънцето е в Риби или Овен", text)

    def test_only_the_person_without_a_time_is_named(self):
        text = birthtime.note([("Иван", self.known), ("Мария", self.unknown)])
        self.assertIn("За Мария", text)
        self.assertNotIn("Иван", text)

    def test_the_prompt_rule_names_the_people_and_overrides_the_templates(self):
        rule = birthtime.block(["Мария"])
        self.assertIn("Мария", rule)
        self.assertIn("overrides every earlier instruction", rule)
        self.assertIn("NO Ascendant", rule)


class PromptTest(unittest.TestCase):
    """Правилото за неизвестен час влиза в системния промпт на всяка AI заявка и се маха след анализа."""

    def system_prompt(self, names=None, add_context=True):
        captured = {}

        class FakeResponse:
            status_code = 200

            def json(self):
                return {"choices": [{"message": {"content": "отговор"}, "finish_reason": "stop"}], "usage": {}}

        class FakeClient:
            def __init__(self, *a, **k):
                pass

            async def __aenter__(self):
                return self

            async def __aexit__(self, *a):
                return False

            async def post(self, url, headers=None, json=None, **k):
                captured["system"] = json["messages"][0]["content"]
                captured["user"] = json["messages"][1]["content"]
                return FakeResponse()

        import asyncio

        async def run():
            token = birthtime.bind(names) if names is not None else None
            try:
                await main.ai_interpreter._call_api("система", "потребител", 100, add_context=add_context)
            finally:
                if token is not None:
                    birthtime.unbind(token)

        ai = main.ai_interpreter
        with mock.patch.object(ai, "ollama_key", "k"), mock.patch.object(ai, "ollama_url", "http://x"), \
                mock.patch("ai_interpreter.httpx.AsyncClient", FakeClient):
            asyncio.run(run())
        self.last_user = captured["user"]
        return captured["system"]

    def test_the_reminder_ends_the_user_prompt(self):
        self.system_prompt(["Мария"])
        self.assertTrue(self.last_user.rstrip().endswith("do not write those parts."))
        self.assertIn("Мария", self.last_user)
        self.system_prompt()
        self.assertNotIn("REMINDER", self.last_user)

    def test_the_rule_is_the_last_part_of_the_system_prompt(self):
        prompt = self.system_prompt(["Мария"])
        self.assertIn("BIRTH TIME UNKNOWN", prompt)
        self.assertIn("ПРАВИЛА ЗА БЕЗОПАСНОСТ", prompt)
        self.assertGreater(prompt.index("BIRTH TIME UNKNOWN"), prompt.index("ПРАВИЛА ЗА БЕЗОПАСНОСТ"))

    def test_no_rule_without_a_binding_or_for_technical_requests(self):
        self.assertNotIn("BIRTH TIME UNKNOWN", self.system_prompt())
        self.assertNotIn("BIRTH TIME UNKNOWN", self.system_prompt(["Мария"], add_context=False))

    def test_the_binding_does_not_leak(self):
        self.system_prompt(["Мария"])
        self.assertEqual(birthtime.current.get(), "")
        self.assertEqual(birthtime.reminder.get(), "")


class CheckerTest(unittest.TestCase):
    """Текстът не може да твърди домове, Асцендент или MC за човек без час, а знакът на Луната е един от възможните."""

    @classmethod
    def setUpClass(cls):
        eng = engine.get_engine()
        cls.unknown = eng.calculate_chart("1990-05-15", "", SOFIA[0], SOFIA[1], time_known=False)
        cls.equinox = eng.calculate_chart("2026-03-20", "", SOFIA[0], SOFIA[1], time_known=False)
        cls.known = eng.calculate_chart("1992-07-20", "09:30", 43.2141, 27.9147)

    def single(self, chart=None):
        return text_check.build_facts(mode="natal", user_name="Мария", natal_chart=chart or self.unknown)

    def pair(self):
        return text_check.build_facts(mode="natal", user_name="Мария", natal_chart=self.unknown,
                                      partner_name="Иван", partner_chart=self.known)

    def codes(self, text, facts):
        return [v.code for v in text_check.check_text(text, facts)]

    def test_facts_know_who_has_no_houses(self):
        facts = self.single()
        self.assertEqual(facts.no_houses, {"user"})
        self.assertEqual(facts.cusps["user"], {})
        self.assertEqual(facts.alt_signs[("user", "Moon")], ["Capricorn", "Aquarius"])
        self.assertEqual(self.pair().no_houses, {"user"})

    def test_a_house_claim_for_a_person_without_a_time_is_refused(self):
        facts = self.single()
        for text in ("Слънцето е в 7-ми дом.", "Венера е в десетия дом.", "Асцендентът е в Лъв.", "MC е в Телец.",
                     "Седмият дом е управляван от Венера.", "Шестият дом започва в Рак."):
            with self.subTest(text=text):
                self.assertIn("no_time", self.codes(text, facts))

    def test_a_sentence_that_explains_the_missing_time_is_not_a_claim(self):
        facts = self.single()
        for text in ("Часът на раждане не е известен, затова няма Асцендент и домове.",
                     "Без час на раждане не можем да определим Асцендента.",
                     "Непознатият час не позволява да се покаже домът на Слънцето."):
            with self.subTest(text=text):
                self.assertNotIn("no_time", self.codes(text, facts))

    def test_planets_in_signs_and_aspects_are_not_affected(self):
        facts = self.single()
        sun = self.unknown["planets"]["Sun"]["zodiac_sign"]
        bg = text_check.SIGN_BG[sun]
        self.assertEqual(self.codes(f"Слънцето е в {bg}.", facts), [])

    def test_the_moon_sign_must_be_one_of_the_possible_ones(self):
        facts = self.single()
        self.assertEqual(self.codes("Луната е в Козирог или в Водолей, според часа.", facts), [])
        self.assertEqual(self.codes("Луната е в Козирог.", facts), [])
        self.assertIn("sign", self.codes("Луната е в Телец.", facts))

    def test_a_body_that_changes_sign_may_be_in_either_sign(self):
        facts = self.single(self.equinox)
        self.assertEqual(self.codes("Слънцето е в Риби.", facts), [])
        self.assertEqual(self.codes("Слънцето е в Овен.", facts), [])
        self.assertIn("sign", self.codes("Слънцето е в Лъв.", facts))

    def test_with_two_people_only_the_one_without_a_time_is_refused(self):
        facts = self.pair()
        venus_house = self.known["planets"]["Venus"]["house"]
        ordinal = {1: "първия", 2: "втория", 3: "третия", 4: "четвъртия", 5: "петия", 6: "шестия", 7: "седмия",
                   8: "осмия", 9: "деветия", 10: "десетия", 11: "единадесетия", 12: "дванадесетия"}[venus_house]
        self.assertNotIn("no_time", self.codes(f"Венера на Иван е в {ordinal} дом.", facts))
        self.assertIn("no_time", self.codes("Венера на Мария е в 7-ми дом.", facts))
        self.assertIn("no_time", self.codes("Асцендентът на Мария е в Лъв.", facts))

    def test_nothing_changes_when_every_time_is_known(self):
        facts = text_check.build_facts(mode="natal", user_name="Иван", natal_chart=self.known)
        self.assertEqual(facts.no_houses, set())
        house = self.known["planets"]["Sun"]["house"]
        self.assertNotIn("no_time", self.codes(f"Слънцето е в {house}-ти дом.", facts))


class JobTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(main.app)

    def setUp(self):
        limiter.reset()

    def post(self, headers, body, **params):
        query = "&".join(f"{k}={v}" for k, v in {"wait": 60, **params}.items())
        return self.client.post(f"/jobs?{query}", json=body, headers=headers)

    def spy(self):
        seen = []

        async def fake(*args, **kwargs):
            seen.append({**kwargs, "rule": birthtime.current.get()})
            return "<p>текст</p>"
        return seen, mock.patch.object(main.ai_interpreter, "interpret_chart", fake)

    NO_TIME = {"date": "1990-05-15", "lat": 42.6977, "lon": 23.3219, "birth_time_known": False, "name": "Мария"}
    PAIR = {"partner_name": "Иван", "partner_date": "1992-07-20", "partner_time": "09:30", "partner_lat": 43.2141,
            "partner_lon": 27.9147}

    def test_a_single_analysis_without_a_time(self):
        h = register_and_login(self.client, "ut-single@test.bg")
        seen, patch = self.spy()
        with patch:
            job = self.post(h, self.NO_TIME).json()["job"]
        self.assertEqual(job["status"], "succeeded", job)
        chart = seen[0]["natal_chart"]
        self.assertIs(chart["time_known"], False)
        self.assertEqual((chart["houses"], chart["angles"]), ({}, {}))
        self.assertIn("BIRTH TIME UNKNOWN", seen[0]["rule"])
        self.assertIn("Мария", seen[0]["rule"])
        self.assertEqual(birthtime.current.get(), "")

    def test_the_saved_report_starts_with_the_note_and_keeps_no_made_up_time(self):
        h = register_and_login(self.client, "ut-saved@test.bg")
        seen, patch = self.spy()
        with patch:
            job = self.post(h, self.NO_TIME).json()["job"]
        report = self.client.get(f"/reports/{job['report_id']}", headers=h).json()
        self.assertTrue(report["content"].startswith("**Бележка за часа на раждане.**"), report["content"][:120])
        self.assertIn("Козирог или Водолей", report["content"])
        self.assertIn("<p>текст</p>", report["content"])
        params = report["params"]
        self.assertIs(params["birth_time_known"], False)
        self.assertNotIn("time", params)

    def test_a_known_time_has_no_note_and_no_rule(self):
        h = register_and_login(self.client, "ut-known@test.bg")
        seen, patch = self.spy()
        with patch:
            job = self.post(h, CHART).json()["job"]
        self.assertEqual(seen[0]["rule"], "")
        content = self.client.get(f"/reports/{job['report_id']}", headers=h).json()
        self.assertNotIn("Бележка за часа на раждане", content["content"])
        self.assertNotIn("birth_time_known", content["params"])

    def test_only_the_second_person_without_a_time(self):
        h = register_and_login(self.client, "ut-partner@test.bg")
        seen, patch = self.spy()
        body = {**CHART, **self.PAIR, "partner_date": "1990-05-15", "partner_time_known": False, "partner_time": None,
                "partner_name": "Мария"}
        with patch:
            job = self.post(h, body).json()["job"]
        self.assertEqual(job["status"], "succeeded", job)
        self.assertEqual(len(seen[0]["natal_chart"]["houses"]), 12)
        self.assertEqual(seen[0]["partner_chart"]["houses"], {})
        self.assertIn("Мария", seen[0]["rule"])
        content = self.client.get(f"/reports/{job['report_id']}", headers=h).json()
        self.assertIs(content["params"]["partner_time_known"], False)
        self.assertNotIn("partner_time", content["params"])
        self.assertIn("За Мария", content["content"])

    def test_a_time_is_still_required_when_not_marked_unknown(self):
        h = register_and_login(self.client, "ut-required@test.bg")
        r = self.post(h, {"date": "1990-05-15", "lat": 42.6977, "lon": 23.3219})
        self.assertEqual(r.status_code, 422, r.text)
        self.assertEqual(self.client.get("/jobs", headers=h).json()["jobs"], [])

    def test_a_period_forecast_without_a_time(self):
        h = register_and_login(self.client, "ut-period@test.bg")
        events = []

        async def month(*args, **kwargs):
            events.append(kwargs["monthly_events"])
            assert kwargs["natal_chart"]["houses"] == {}
            return "<p>месец</p>"
        overview = mock.AsyncMock(return_value="<p>преглед</p>")
        body = {**self.NO_TIME, "is_dynamic": True, "target_date": "2026-01-01", "end_date": "2026-01-31"}
        with mock.patch.object(main.ai_interpreter, "_process_monthly_chunk", month), \
                mock.patch.object(main.ai_interpreter, "compose_period_overview", overview):
            job = self.post(h, body).json()["job"]
        self.assertEqual(job["status"], "succeeded", job)
        rows = [row for month_rows in events for row in month_rows if row.get("type") == "TRANSIT"]
        self.assertTrue(rows)
        for row in rows:
            self.assertNotIn("transit_planet_natal_house", row)
            self.assertNotIn("natal_planet_natal_house", row)
            self.assertNotEqual(row["natal_planet"], "Moon")
        content = self.client.get(f"/reports/{job['report_id']}", headers=h).json()["content"]
        self.assertIn("Бележка за часа на раждане", content)
        self.assertLess(content.index("Бележка за часа на раждане"), content.index("<p>преглед</p>"))

    def test_the_calculate_endpoint(self):
        body = {"date": "1990-05-15", "lat": 42.6977, "lon": 23.3219, "birth_time_known": False}
        r = self.client.post("/calculate", json=body)
        self.assertEqual(r.status_code, 200, r.text)
        data = r.json()
        self.assertIs(data["time_known"], False)
        self.assertEqual((data["houses"], data["angles"]), ({}, {}))
        self.assertIn("Moon", data["sign_ranges"])
        known = self.client.post("/calculate", json={**CHART, "time": "14:30"}).json()
        self.assertIs(known["time_known"], True)
        self.assertEqual(len(known["houses"]), 12)
        gap = self.client.post("/calculate", json={"date": "2026-03-29", "time": "03:30", "lat": 42.6977, "lon": 23.3219})
        self.assertEqual(gap.status_code, 400)
        self.assertIn("не съществува", gap.json()["detail"])


class DocxTest(unittest.TestCase):
    """Изнасянето не показва домове и Асцендент без час и не се чупи от Луната без позиция (подробно: test_report_export.py)."""

    def text_of(self, chart, aspects):
        import io
        from docx import Document
        from docx_generator import DOCXGenerator
        data = {"title": "Анализ", "kind": "Общ анализ", "created": "10.10.2026", "people": [],
                "charts": [{"title": "Карта: Мария", "chart": chart, "aspects": aspects}],
                "sections": [{"title": "", "text": "Текст."}]}
        doc = Document(io.BytesIO(DOCXGenerator().generate_docx(data)))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts += [cell.text for cell in row.cells]
        return "\n".join(parts)

    def test_without_a_time(self):
        eng = engine.get_engine()
        chart = eng.calculate_chart("1990-05-15", "", SOFIA[0], SOFIA[1], time_known=False)
        text = self.text_of(chart, aspects_engine.calculate_natal_aspects(chart))
        self.assertNotIn("\nДомове", text)
        self.assertNotIn("\nАсцендент\n", text)
        self.assertNotIn("празен", text)
        self.assertIn("Луна\nКозирог или Водолей (според часа)", text)
        self.assertIn("Аспекти", text)

    def test_with_a_time_nothing_is_lost(self):
        eng = engine.get_engine()
        chart = eng.calculate_chart("1990-05-15", "14:30", SOFIA[0], SOFIA[1])
        text = self.text_of(chart, aspects_engine.calculate_natal_aspects(chart))
        self.assertIn("Домове", text)
        self.assertIn("\nАсцендент\n", text)
        self.assertIn("Аспекти", text)


class RealPromptTest(unittest.TestCase):
    """Истинските шаблони и блокове с подменен само вик към AI: без час в подканата няма домове, куспиди и Асцендент."""

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(main.app)

    def setUp(self):
        limiter.reset()

    def run_job(self, email, body):
        h = register_and_login(self.client, email)
        calls = []

        async def record(system_prompt, user_prompt, max_tokens, **kwargs):
            calls.append({"system": system_prompt, "user": user_prompt, "rule": birthtime.current.get(),
                          "reminder": birthtime.reminder.get()})
            if "MODE: PERIOD OVERVIEW" in system_prompt:
                return "**Накратко**: преглед."
            return "<p>Подменен текст.</p>"
        with mock.patch.object(main.ai_interpreter, "_call_api", record):
            response = self.client.post("/jobs?wait=60", json=body, headers=h)
        self.assertEqual(response.status_code, 200, response.text)
        job = response.json()["job"]
        self.assertEqual(job["status"], "succeeded", job)
        self.assertTrue(calls)
        return calls

    def assert_no_house_data(self, calls):
        for call in calls:
            text = call["user"]
            for marker in ('"house":', '"houses"', "Ascendant_formatted", '"cusp"', "Ascendant_sign",
                           "NATAL HOUSES (CALCULATED) ---"):
                self.assertFalse(marker in text, marker)             # без assertNotIn: той печата целия промпт
            self.assertIn("BIRTH TIME UNKNOWN", call["rule"])
            self.assertTrue(call["reminder"].startswith("REMINDER (birth time unknown for:"))

    def test_single_natal(self):
        self.assert_no_house_data(self.run_job("rp-natal@test.bg", {
            "date": "1990-05-15", "lat": 42.6977, "lon": 23.3219, "birth_time_known": False, "name": "Мария"}))

    def test_single_with_a_date_for_transits(self):
        self.assert_no_house_data(self.run_job("rp-transit@test.bg", {
            "date": "1990-05-15", "lat": 42.6977, "lon": 23.3219, "birth_time_known": False, "name": "Мария",
            "target_date": "2026-10-10", "target_time": "12:00"}))

    def test_two_people_without_a_time_with_a_date_for_transits(self):
        self.assert_no_house_data(self.run_job("rp-pair-transit@test.bg", {
            "date": "1990-05-15", "lat": 42.6977, "lon": 23.3219, "birth_time_known": False, "name": "Мария",
            "partner_name": "Ана", "partner_date": "1992-07-20", "partner_lat": 43.2141, "partner_lon": 27.9147,
            "partner_time_known": False, "target_date": "2026-10-10", "target_time": "12:00"}))

    def test_period_forecast_for_two_people_without_a_time(self):
        calls = self.run_job("rp-period@test.bg", {
            "date": "1990-05-15", "lat": 42.6977, "lon": 23.3219, "birth_time_known": False, "name": "Мария",
            "partner_name": "Ана", "partner_date": "1992-07-20", "partner_lat": 43.2141, "partner_lon": 27.9147,
            "partner_time_known": False, "is_dynamic": True, "target_date": "2026-01-01", "end_date": "2026-01-31"})
        self.assertGreaterEqual(len(calls), 2)               # месец и общ преглед
        self.assert_no_house_data(calls)

    def test_one_person_with_a_time_keeps_the_houses_in_the_prompt(self):
        calls = self.run_job("rp-mixed@test.bg", {
            "date": "1990-05-15", "lat": 42.6977, "lon": 23.3219, "time": "14:30", "name": "Иван",
            "partner_name": "Мария", "partner_date": "1992-07-20", "partner_lat": 43.2141, "partner_lon": 27.9147,
            "partner_time_known": False})
        user = calls[0]["user"]
        self.assertTrue('"houses"' in user)                                   # Иван има домове
        self.assertTrue("--- PARTNER PLANETS IN USER'S NATAL HOUSES (CALCULATED) ---" in user)
        self.assertFalse("--- ИВАН PLANETS IN МАРИЯ'S NATAL HOUSES (CALCULATED) ---" in user)   # Мария няма домове
        self.assertTrue(calls[0]["reminder"].startswith("REMINDER (birth time unknown for: Мария)"))
        self.assertIn("BIRTH TIME UNKNOWN", calls[0]["rule"])


if __name__ == "__main__":
    unittest.main()
